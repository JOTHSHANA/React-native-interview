from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parents[1] / "ReqRes_App_Deliverables.docx"
GITHUB_URL = "https://github.com/JOTHSHANA/React-native-interview"


def set_run_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), kwargs[edge].get("val", "single"))
            element.set(qn("w:sz"), kwargs[edge].get("sz", "4"))
            element.set(qn("w:color"), kwargs[edge].get("color", "5B4BFF"))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def add_heading_custom(doc, text, size=22):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=(26, 21, 48))
    p.paragraph_format.space_after = Pt(6)
    return p


def add_body(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=11, bold=bold, color=(26, 21, 48))
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=11, color=(26, 21, 48))
    p.paragraph_format.space_after = Pt(3)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, name="Consolas", size=10, color=(44, 31, 168))
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.2)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("React Native API Integration Task")
    set_run_font(run, size=26, bold=True, color=(91, 75, 255))
    title.paragraph_format.space_after = Pt(2)

    subtitle = doc.add_paragraph()
    run = subtitle.add_run("Deliverables document  ·  ReqRes App")
    set_run_font(run, size=13, bold=True, color=(107, 102, 128))
    subtitle.paragraph_format.space_after = Pt(16)

    # GitHub source box
    table = doc.add_table(rows=2, cols=1)
    table.autofit = True
    hdr = table.cell(0, 0)
    body = table.cell(1, 0)
    shade_cell(hdr, "5B4BFF")
    shade_cell(body, "EEEDFF")
    p = hdr.paragraphs[0]
    r = p.add_run("SOURCE CODE  —  GitHub repository")
    set_run_font(r, size=11, bold=True, color=(255, 255, 255))
    p = body.paragraphs[0]
    r = p.add_run(GITHUB_URL)
    set_run_font(r, size=12, bold=True, color=(44, 31, 168))
    p2 = body.add_paragraph()
    r = p2.add_run(
        "Clone this repository. Copy .env.example to .env and add the ReqRes API key. Do not commit .env."
    )
    set_run_font(r, size=10, color=(107, 102, 128))
    doc.add_paragraph()

    add_heading_custom(doc, "1. Project overview", 16)
    add_body(
        doc,
        "A React Native (Expo) mobile app that integrates with the ReqRes REST API. It covers authentication, form validation, a centralized API layer, profile and paginated products, loading/error/retry states, and React Navigation.",
    )
    add_body(doc, "Base URL: https://reqres.in", bold=True)
    add_body(
        doc,
        "Flow: Launch app → Login / Register → Home → Profile / Products. After login, a bottom tab bar is used for Home, Products, Profile, and Logout.",
    )

    add_heading_custom(doc, "2. Tech stack", 16)
    stack = [
        "React Native with Expo SDK 57 (JavaScript)",
        "React Navigation (native stack + bottom tabs)",
        "Axios centralized API client with x-api-key on every request",
        "Context API for auth state (token, login, register, logout)",
        "AsyncStorage for token persistence",
        "Environment variables via .env (EXPO_PUBLIC_API_KEY)",
    ]
    for item in stack:
        add_bullet(doc, item)

    add_heading_custom(doc, "3. Deliverables", 16)

    add_heading_custom(doc, "Deliverable 1 — Complete React Native source code", 13)
    add_body(
        doc,
        "The full Expo project is on GitHub. It includes App.js, index.js, package.json, app.json, .env.example, README.md, and the src/ folder. Clone the repo, run npm install, then npx expo start. The real API key stays in a local .env file and is not in the repository.",
    )
    add_code(
        doc,
        "src/\n"
        "  api/           apiClient, authApi, userApi, productApi\n"
        "  screens/       Login, Register, Home, Profile, Products\n"
        "  components/    Button, Input, Loader, ProductCard, TabBar\n"
        "  navigation/    Auth stack and tab navigator\n"
        "  constants/     API URLs and theme\n"
        "  utils/         Form validation\n"
        "  storage/       Token persistence\n"
        "  context/       Auth state",
    )

    add_heading_custom(doc, "Deliverable 2 — All 5 screens implemented", 13)
    screens = [
        "Login — email, password, Continue button, register link, error banner, loading spinner. Test values are prefilled.",
        "Register — email, password, confirm password, Create account button, login link, client-side validation.",
        "Home — greeting, live API badge, hero card, product stats, shortcuts to Products and Profile.",
        "Profile — avatar, first name, last name, email, user ID, loading/error/retry, Logout button.",
        "Products — FlatList/grid of color products (name, year, color, pantone), pull-to-refresh, page 1 then Load more for page 2.",
    ]
    for item in screens:
        add_bullet(doc, item)

    add_heading_custom(doc, "Deliverable 3 — API service layer", 13)
    add_body(
        doc,
        "All HTTP calls go through a shared Axios client (src/api/apiClient.js) with a 10-second timeout, JSON headers, and x-api-key. URLs live in src/constants/apiConstants.js. Duplicate in-flight requests are ignored.",
    )
    apis = [
        "POST /api/login — loginUser in authApi.js",
        "POST /api/register — registerUser in authApi.js (user registration only, not orders)",
        "GET /api/users/2 — getUserDetails in userApi.js",
        "GET /api/products?page=n — getProducts in productApi.js",
    ]
    for item in apis:
        add_bullet(doc, item)

    add_heading_custom(doc, "Deliverable 4 — Navigation setup", 13)
    add_body(
        doc,
        "React Navigation is used throughout. If no token is stored, the user sees the Auth stack (Login, Register). After a successful login or register, the token is saved and the app switches to tabs: Home, Products, Profile, plus Logout (opens a confirm dialog). Logout also exists on the Profile screen.",
    )

    add_heading_custom(doc, "Deliverable 5 — Form validation", 13)
    add_body(doc, "Validation lives in src/utils/validation.js and runs before any login/register API call.")
    rules = [
        "Email: required and must be a valid format",
        "Password: required, minimum 6 characters",
        "Confirm password: required and must match password",
        "Invalid forms do not call the API",
        "Submit buttons are disabled and show a spinner while a request is in flight",
    ]
    for item in rules:
        add_bullet(doc, item)

    add_heading_custom(doc, "Deliverable 6 — Error handling", 13)
    errors = [
        "Network errors and timeouts mapped to user-friendly messages",
        "Invalid credentials and registration failures shown on Login/Register",
        "Profile and Products show error text plus a Retry button",
        "Products pagination errors appear above Load more",
        "Empty product list has an empty state",
        "App-level ErrorBoundary in App.js with a Try again action",
        "Logout confirm dialog works on both phone (Expo Go) and web",
    ]
    for item in errors:
        add_bullet(doc, item)

    add_heading_custom(doc, "Deliverable 7 — README with setup instructions", 13)
    add_body(
        doc,
        "README.md in the project root explains npm install, .env setup, Expo Go, tunnel mode, test accounts, and folder structure.",
    )

    add_heading_custom(doc, "4. How to run", 16)
    add_body(doc, "Clone from GitHub")
    add_code(
        doc,
        "git clone https://github.com/JOTHSHANA/React-native-interview.git\n"
        "cd React-native-interview\nnpm install\n"
        "# copy .env.example to .env and add EXPO_PUBLIC_API_KEY",
    )
    add_body(doc, "Phone (Expo Go)")
    add_code(doc, "npx expo start\n# or: npx expo start --tunnel")
    add_body(doc, "Laptop browser")
    add_code(doc, "npx expo start --web")

    add_heading_custom(doc, "5. Test credentials", 16)
    creds = doc.add_table(rows=3, cols=3)
    creds.style = "Table Grid"
    headers = ["Action", "Email", "Password"]
    values = [
        ["Login", "eve.holt@reqres.in", "cityslicka"],
        ["Register", "eve.holt@reqres.in", "pistol"],
    ]
    for i, h in enumerate(headers):
        cell = creds.rows[0].cells[i]
        shade_cell(cell, "5B4BFF")
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=11, bold=True, color=(255, 255, 255))
    for r, row in enumerate(values, start=1):
        for c, val in enumerate(row):
            cell = creds.rows[r].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=11, color=(26, 21, 48))
    doc.add_paragraph()
    add_body(
        doc,
        "ReqRes only accepts these defined users for a successful login or register.",
    )

    add_heading_custom(doc, "6. Important notes covered", 16)
    notes = [
        "API key is not hardcoded in source; it is loaded from .env",
        "React Navigation is used for all screen movement",
        "Auth state uses Context API (not Redux)",
        "ErrorBoundary wraps the root app",
        "App is tested with the credentials above",
        "Register calls POST /api/register for user signup, not order placement",
    ]
    for item in notes:
        add_bullet(doc, item)

    add_heading_custom(doc, "7. Source code location", 16)
    add_body(
        doc,
        "Source code is hosted on GitHub (no zip file): https://github.com/JOTHSHANA/React-native-interview",
    )

    footer = doc.add_paragraph()
    run = footer.add_run(
        "ReqRes React Native API Integration  ·  Deliverables summary"
    )
    set_run_font(run, size=9, color=(107, 102, 128))

    try:
        doc.save(OUTPUT)
        print(OUTPUT)
    except PermissionError:
        fallback = OUTPUT.with_name("ReqRes_App_Deliverables_github.docx")
        doc.save(fallback)
        print(fallback)


if __name__ == "__main__":
    main()
